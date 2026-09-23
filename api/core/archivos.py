import io
from datetime import date

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

TURQUESA = (0, 153, 179)

TITULOS = {
    "id_cobro": "ID", "fecha_cobro": "Fecha cobro", "fecha_vencimiento": "Vencimiento",
    "fecha_pago": "Fecha pago", "donante": "Donante", "linea_estrategica": "Línea estratégica",
    "asignacion": "Asignación", "campana": "Campaña", "forma_pago": "Forma de pago",
    "importe_comprometido": "Comprometido", "importe_cobrado": "Cobrado", "estatus": "Estatus",
    "vencido": "Vencido", "dias_atraso": "Días de atraso",
    "id_llamada": "ID", "fecha_llamada": "Fecha", "telefonista": "Telefonista",
    "resultado": "Resultado", "comentarios": "Comentarios",
    "id_meta": "ID", "desde": "Desde", "hasta": "Hasta", "objetivo_total": "Objetivo total",
    "objetivo_periodo": "Objetivo del periodo", "comprometido": "Comprometido",
    "cobrado": "Cobrado", "faltante": "Faltante", "porcentaje_avance": "Avance %",
}

MONTOS = {"importe_comprometido", "importe_cobrado", "objetivo_total", "objetivo_periodo",
          "comprometido", "cobrado", "faltante"}
ENTEROS = {"id_cobro", "id_llamada", "id_meta", "dias_atraso", "porcentaje_avance"}
FECHAS = {"fecha_cobro", "fecha_vencimiento", "fecha_pago", "desde", "hasta"}



def titulo_columna(columna: str) -> str:
    return TITULOS.get(columna, columna.replace("_", " ").capitalize())


def valor_legible(columna: str, valor) -> str:
    if columna in MONTOS:
        return f"${float(valor or 0):,.2f}"
    if columna == "porcentaje_avance":
        return f"{valor} %"
    return "" if valor is None else str(valor)


def latin1(texto: str) -> str:
    texto = texto.replace("–", "-").replace("—", "-").replace("…", "...")
    return texto.encode("latin-1", "replace").decode("latin-1")


def anchos_de_columnas(pdf: FPDF, encabezados: list, filas: list) -> list:
    minimos, preferidos = [], []
    for indice, encabezado in enumerate(encabezados):
        pdf.set_font("Helvetica", "B", 7)
        palabra_mas_larga = max(pdf.get_string_width(p) for p in encabezado.split())
        pdf.set_font("Helvetica", "", 7)
        textos = [fila[indice] for fila in filas]
        palabra_mas_larga = max([palabra_mas_larga] + [pdf.get_string_width(p) for t in textos for p in t.split()])
        texto_mas_largo = max([pdf.get_string_width(t) for t in textos] + [palabra_mas_larga])
        minimos.append(palabra_mas_larga + 3)
        preferidos.append(min(max(texto_mas_largo + 3, palabra_mas_larga + 3), 60))

    disponible = pdf.epw
    if sum(preferidos) <= disponible:
        return [p * disponible / sum(preferidos) for p in preferidos]
    if sum(minimos) >= disponible:
        return [m * disponible / sum(minimos) for m in minimos]
    factor = (disponible - sum(minimos)) / (sum(preferidos) - sum(minimos))
    return [m + (p - m) * factor for m, p in zip(minimos, preferidos)]


class DocumentoPDF(FPDF):
    def footer(self):
        self.set_y(-10)
        self.set_font("Helvetica", "", 7)
        self.set_text_color(120)
        self.cell(0, 5, latin1(f"Cáritas de Monterrey, A.B.P. · Sistema de Ingresos · Página {self.page_no()}"),
                  align="C")


def archivo_pdf(titulo: str, detalles: list, columnas: list, filas: list) -> bytes:
    pdf = DocumentoPDF(orientation="L", unit="mm", format="Letter")
    pdf.set_auto_page_break(True, margin=14)
    pdf.set_margins(12, 12, 12)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 17)
    pdf.set_text_color(*TURQUESA)
    pdf.cell(0, 9, latin1(titulo), new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(60)
    for etiqueta, valor in detalles:
        pdf.cell(0, 5, latin1(f"{etiqueta}: {valor}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_text_color(30)
    if not filas:
        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 8, latin1("Sin registros en este periodo."), new_x="LMARGIN", new_y="NEXT")
        return bytes(pdf.output())

    encabezados = [latin1(titulo_columna(c)) for c in columnas]
    textos = [[latin1(valor_legible(c, v)) for c, v in zip(columnas, fila)] for fila in filas]
    anchos = anchos_de_columnas(pdf, encabezados, textos)
    alineaciones = ["RIGHT" if c in MONTOS or c in ENTEROS else "LEFT" for c in columnas]
    encabezado = FontFace(emphasis="BOLD", color=255, fill_color=TURQUESA)

    pdf.set_font("Helvetica", "", 7)
    with pdf.table(col_widths=anchos, text_align=alineaciones, headings_style=encabezado,
                   line_height=4.5, cell_fill_color=(240, 244, 247), cell_fill_mode="ROWS") as tabla:
        tabla.row(encabezados)
        for fila in textos:
            tabla.row(fila)

    return bytes(pdf.output())


def valor_excel(columna: str, valor):
    if valor in (None, ""):
        return None
    if columna in MONTOS:
        return float(valor)
    if columna in ENTEROS:
        return int(valor)
    if columna in FECHAS:
        return date.fromisoformat(str(valor))
    return valor


def archivo_excel(titulo: str, detalles: list, columnas: list, filas: list) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Datos"

    hoja.append([titulo_columna(c) for c in columnas])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0099B3")
        celda.alignment = Alignment(horizontal="center")

    for fila in filas:
        hoja.append([valor_excel(c, v) for c, v in zip(columnas, fila)])

    for indice, columna in enumerate(columnas, start=1):
        letra = hoja.cell(row=1, column=indice).column_letter
        largos = [len(titulo_columna(columna))] + [len(valor_legible(columna, f[indice - 1])) for f in filas]
        hoja.column_dimensions[letra].width = min(max(largos) + 3, 45)
        formato = ('"$"#,##0.00' if columna in MONTOS else
                   "yyyy-mm-dd" if columna in FECHAS else None)
        if formato:
            for (celda,) in hoja.iter_rows(min_row=2, min_col=indice, max_col=indice):
                celda.number_format = formato

    hoja.freeze_panes = "A2"
    if filas:
        hoja.auto_filter.ref = hoja.dimensions

    resumen = libro.create_sheet("Resumen", 0)
    resumen.append([titulo])
    resumen["A1"].font = Font(bold=True, size=14, color="0099B3")
    resumen.append([])
    for etiqueta, valor in detalles:
        resumen.append([etiqueta, valor])
        resumen.cell(row=resumen.max_row, column=1).font = Font(bold=True)
    resumen.append(["Registros", len(filas)])
    resumen.cell(row=resumen.max_row, column=1).font = Font(bold=True)
    resumen.column_dimensions["A"].width = 22
    resumen.column_dimensions["B"].width = 50

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
